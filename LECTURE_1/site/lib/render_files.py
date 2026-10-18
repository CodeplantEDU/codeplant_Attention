"""Render model-authored data to files. Never execute model code or fetch URLs."""
import csv
import io
import os
import json
import re
import sys
import textwrap
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

ASSETS = Path(os.environ.get('DIALOGUE_ASSETS') or Path(__file__).resolve().parents[1] / 'public' / 'assets')


def plain(text):
    return re.sub(r'\*\*(.*?)\*\*|`([^`]+)`', lambda m: m[1] or m[2], text)


def blocks(content):
    lines = content.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('|') and i + 1 < len(lines) and re.fullmatch(r'[| :\-]+', lines[i + 1].strip()):
            rows = [[plain(c.strip()) for c in line.strip('|').split('|')]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([plain(c.strip()) for c in lines[i].strip().strip('|').split('|')]); i += 1
            yield 'table', rows
            continue
        if line:
            heading = re.match(r'^(#{1,6})\s+(.+)', line)
            bullet = re.match(r'^(?:[-*+]\s+|\d+[.)]\s+)(.+)', line)
            if heading:
                yield 'heading', (len(heading[1]), plain(heading[2]))
            elif bullet:
                yield 'bullet', plain(bullet[1])
            elif not line.startswith('```'):
                yield 'paragraph', line
        i += 1


def docx_file(file, target):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    doc.core_properties.author = 'CODEPLANT'
    doc.core_properties.created = doc.core_properties.modified = datetime.now(timezone.utc)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    for name, size in [('Normal', 11), ('Title', 24), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 12)]:
        style = doc.styles[name]
        style.font.name = 'Pretendard'; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Pretendard')
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.35
    if (ASSETS / 'codeplant-logo.png').exists():
        section.header.paragraphs[0].add_run().add_picture(str(ASSETS / 'codeplant-logo.png'), width=Inches(1.15))
    footer = section.footer.paragraphs[0]; footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    title = file.get('title') or Path(file['name']).stem
    doc.add_paragraph(title, 'Title')
    first = True
    for kind, value in blocks(file['content']):
        if kind == 'heading':
            if first and value[1] == title:
                first = False; continue
            doc.add_heading(value[1], level=min(value[0], 3))
        elif kind == 'table':
            width = max(map(len, value))
            if width > 12:
                raise ValueError('Word 표는 12열 이하로 나눠주세요.')
            table = doc.add_table(rows=0, cols=width)
            table.autofit = False
            weights = [max(6,min(42,max(len(row[col]) if col<len(row) else 0 for row in value))) ** .5 for col in range(width)]
            for col, weight in enumerate(weights): table.columns[col].width = Inches(6.5*weight/sum(weights))
            borders = OxmlElement('w:tblBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                element = OxmlElement('w:' + edge); element.set(qn('w:val'), 'single'); element.set(qn('w:sz'), '4'); element.set(qn('w:color'), 'D9D9D9'); borders.append(element)
            table._tbl.tblPr.append(borders)
            for n, row in enumerate(value):
                cells = table.add_row().cells
                for col, cell in enumerate(cells):
                    cell.width = Inches(6.5*weights[col]/sum(weights))
                    cell.text = row[col] if col < len(row) else ''
                    tcpr = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement('w:tcMar')
                    for side in ('top', 'bottom', 'start', 'end'):
                        item = OxmlElement('w:' + side); item.set(qn('w:w'), '110'); item.set(qn('w:type'), 'dxa'); margins.append(item)
                    tcpr.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        for run in p.runs:
                            run.font.size = Pt(10.5); run.bold = n == 0
                    if n == 0:
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'D8F1F0'); tcpr.append(shade)
                if n == 0:
                    repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
        else:
            p = doc.add_paragraph(style='List Bullet' if kind == 'bullet' else 'Normal')
            for token in re.split(r'(\*\*.*?\*\*)', value):
                run = p.add_run(token[2:-2] if token.startswith('**') and token.endswith('**') else token)
                run.bold = token.startswith('**') and token.endswith('**')
        first = False
    doc.save(target)


def pdf_file(file, target):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont('CP', str(ASSETS / 'Pretendard-Regular.ttf')))
    pdfmetrics.registerFont(TTFont('CPBold', str(ASSETS / 'Pretendard-Bold.ttf')))
    normal = ParagraphStyle('body', fontName='CP', fontSize=11, leading=17, spaceAfter=8, wordWrap='CJK')
    heading = ParagraphStyle('heading', parent=normal, fontName='CPBold', fontSize=15, leading=22, spaceBefore=12, keepWithNext=True)
    title = ParagraphStyle('title', parent=heading, fontSize=24, leading=32, spaceAfter=20)
    content = [Paragraph(escape(file.get('title') or Path(file['name']).stem), title)]
    for kind, value in blocks(file['content']):
        if kind == 'table':
            cols = max(map(len, value))
            rows = [[Paragraph(escape(c), normal) for c in row + [''] * (cols-len(row))] for row in value]
            table = Table(rows, colWidths=[468 / cols] * cols, repeatRows=1)
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#D8F1F0')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
            content.extend([table, Spacer(1,10)])
        else:
            text = value[1] if kind == 'heading' else ('• ' if kind == 'bullet' else '') + plain(value)
            content.append(Paragraph(escape(text), heading if kind == 'heading' else normal))
    def page(canvas, doc):
        canvas.setFont('CP',9); canvas.drawRightString(540,36,str(doc.page))
    SimpleDocTemplate(str(target), pagesize=letter, rightMargin=72, leftMargin=72, topMargin=60, bottomMargin=54).build(content,onFirstPage=page,onLaterPages=page)


def pptx_file(file, target):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    presentation = Presentation(); presentation.slide_width = Inches(13.333); presentation.slide_height = Inches(7.5)
    title = file.get('title') or Path(file['name']).stem
    sections = []; body = []
    for line in file['content'].splitlines():
        if re.match(r'^#{1,2}\s+', line):
            if body: sections.append((title,body)); body=[]
            title = re.sub(r'^#+\s+', '', line)
        elif line.strip():
            body.extend(textwrap.wrap(plain(line.strip()), width=45))
    if body: sections.append((title,body))
    for title, lines in sections or [(title,[])]:
        for offset in range(0,max(len(lines),1),8):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            if (ASSETS/'codeplant-logo.png').exists(): slide.shapes.add_picture(str(ASSETS/'codeplant-logo.png'),Inches(.65),Inches(.3),width=Inches(1.4))
            box=slide.shapes.add_textbox(Inches(.7),Inches(1),Inches(11.9),Inches(1)).text_frame
            box.word_wrap=True; p=box.paragraphs[0];p.text=title;p.font.name='Pretendard';p.font.size=Pt(30);p.font.bold=True;p.font.color.rgb=RGBColor(45,45,45)
            frame=slide.shapes.add_textbox(Inches(.8),Inches(2.1),Inches(11.6),Inches(4.9)).text_frame;frame.word_wrap=True
            for i,line in enumerate(lines[offset:offset+8]):
                p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=line;p.font.name='Pretendard';p.font.size=Pt(22);p.font.color.rgb=RGBColor(45,45,45);p.space_after=Pt(12)
    presentation.save(target)


def xlsx_file(file, target):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    workbook=Workbook();workbook.remove(workbook.active)
    spec=json.loads(file['content']);sheets=spec.get('sheets')
    if not isinstance(sheets,list) or not 1<=len(sheets)<=10: raise ValueError('Excel 시트는 1~10개가 필요합니다.')
    for n,sheet in enumerate(sheets,1):
        name=re.sub(r'[\[\]:*?/\\]','',str(sheet.get('name',f'표{n}')))[:31] or f'표{n}'
        ws=workbook.create_sheet(name);rows=sheet.get('rows',[])
        if not isinstance(rows,list) or len(rows)>10000: raise ValueError('Excel 행 수를 확인하세요.')
        for i,row in enumerate(rows,1):
            if not isinstance(row,list) or len(row)>100: raise ValueError('Excel 표는 행마다 최대100열입니다.')
            for j,value in enumerate(row,1):
                if not isinstance(value,(str,int,float,bool,type(None))): raise ValueError('Excel 셀은 문자나 숫자여야 합니다.')
                cell=ws.cell(i,j,value)
                if isinstance(value,str): cell.data_type='s'
                cell.font=Font(name='Pretendard',size=11,bold=i==1)
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if i==1: cell.fill=PatternFill('solid',fgColor='D8F1F0')
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(50,max(14,max(len(str(c.value or '')) for c in col)+3))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    workbook.save(target)


def render(file, target):
    ext=file['ext']
    if ext=='docx': docx_file(file,target)
    elif ext=='pdf': pdf_file(file,target)
    elif ext=='pptx': pptx_file(file,target)
    elif ext=='xlsx': xlsx_file(file,target)
    else:
        content=file['content']
        if ext=='csv':
            output=io.StringIO();writer=csv.writer(output)
            for row in csv.reader(io.StringIO(content)):
                writer.writerow(["'"+c if c.startswith(('=','+','@')) else c for c in row])
            content=output.getvalue()
        target.write_text(content,encoding='utf-8-sig' if ext=='csv' else 'utf-8')
    if ext in ('docx','pptx','xlsx'):
        with zipfile.ZipFile(target) as archive:
            if archive.testzip(): raise ValueError('Office 파일 검증에 실패했습니다.')


if __name__=='__main__':
    try:
        packet=json.loads(Path(sys.argv[1]).read_text('utf-8'));folder=Path(sys.argv[2])
        for file in packet['files']:
            if not re.fullmatch(r'[a-f0-9-]{36}',file['id']) or not re.fullmatch(r'[a-z0-9]+',file['ext']): raise ValueError('파일 경로를 확인하세요.')
            render(file,folder/(file['id']+'.'+file['ext']+'.part'))
        print(json.dumps({'ok':True}))
    except ImportError:
        print(json.dumps({'error':'문서 생성 라이브러리가 필요합니다. 사용하는 Python에서 requirements.txt를 설치하세요.'}));sys.exit(1)
    except Exception as exc:
        print(json.dumps({'error':str(exc) if isinstance(exc,ValueError) else '파일 생성 중 오류가 발생했습니다. 요청 형식과 본문을 확인하세요.'}));sys.exit(1)
