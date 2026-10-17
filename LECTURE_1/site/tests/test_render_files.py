import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'lib'))
from render_files import render


class GeneratedFileTests(unittest.TestCase):
    def setUp(self):
        root=Path(__file__).resolve().parents[2]/'tmp'/'render-tests';root.mkdir(parents=True,exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=root);self.folder=Path(self.temp.name)
        self.content='# 공개 수업안\n\n자료에 근거한 수업이다.\n\n## 시간표\n\n| 활동 | 시간 |\n| --- | --- |\n| 도입 | 10분 |\n| 실습 | 110분 |\n\n- 총 120분\n'

    def tearDown(self):
        self.temp.cleanup()

    def test_docx_has_real_title_body_table_and_black_heading_styles(self):
        from docx import Document
        target=self.folder/'output.docx';render({'name':'수업안.docx','ext':'docx','title':'공개 수업안','content':self.content},target)
        doc=Document(target)
        self.assertEqual(doc.paragraphs[0].style.name,'Title')
        self.assertIn('총 120분','\n'.join(p.text for p in doc.paragraphs))
        self.assertEqual(doc.tables[0].cell(2,1).text,'110분')
        self.assertEqual(str(doc.styles['Heading 1'].font.color.rgb),'000000')

    def test_pdf_and_pptx_contain_real_pages_and_slides(self):
        from pypdf import PdfReader
        from pptx import Presentation
        pdf=self.folder/'output.pdf';render({'name':'수업안.pdf','ext':'pdf','title':'공개 수업안','content':self.content},pdf)
        reader=PdfReader(pdf);self.assertGreater(len(reader.pages),0);self.assertIn('120',reader.pages[0].extract_text())
        ppt=self.folder/'output.pptx';render({'name':'강의.pptx','ext':'pptx','title':'공개 수업안','content':self.content},ppt)
        slides=Presentation(ppt).slides;self.assertGreater(len(slides),0)
        self.assertTrue(any('수업안' in shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame))

    def test_xlsx_and_text_do_not_execute_generated_content(self):
        from openpyxl import load_workbook
        target=self.folder/'output.xlsx'
        render({'name':'시간표.xlsx','ext':'xlsx','content':json.dumps({'sheets':[{'name':'시간표','rows':[['활동','시간'],['실습',110],['수식 문자열','=1+1']]}]})},target)
        book=load_workbook(target);self.assertEqual(book.active['B2'].value,110);self.assertEqual(book.active['B3'].data_type,'s')
        text=self.folder/'output.py';render({'name':'예제.py','ext':'py','content':'raise RuntimeError("must not execute")'},text)
        self.assertIn('must not execute',text.read_text('utf-8'))
        csv=self.folder/'output.csv';render({'name':'표.csv','ext':'csv','content':'항목,값\n식,=1+1'},csv)
        self.assertIn("'=1+1",csv.read_text('utf-8-sig'))


if __name__=='__main__':
    unittest.main()
